"""Helpers shared by sources: dates, page text, PDFs and spreadsheets."""

from __future__ import annotations

import hashlib
import io
import logging
import re
from datetime import date, datetime

from bs4 import BeautifulSoup

log = logging.getLogger(__name__)

MONTHS = ("january|february|march|april|may|june|july|august|september|october|november|december|"
          "jan|feb|mar|apr|jun|jul|aug|sep|sept|oct|nov|dec")
_LONG = re.compile(rf"\b({MONTHS})\.?\s+(\d{{1,2}}),?\s+(\d{{4}})\b", re.I)
_NUMERIC = re.compile(r"\b(\d{1,2})/(\d{1,2})/(\d{2,4})\b")
_ISO = re.compile(r"\b(\d{4})[-/](\d{1,2})[-/](\d{1,2})\b")
# "August 2026" (no day): the first of the month.
_MONTH_YEAR = re.compile(rf"\b({MONTHS})\.?\s+(\d{{4}})\b", re.I)


def _month(name: str) -> int:
    return ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"].index(name[:3].lower()) + 1


def parse_date(value) -> date | None:
    """First date found in a string (or a date/datetime), else None."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value)
    found: list[tuple[int, date]] = []
    for m in _ISO.finditer(s):
        try:
            found.append((m.start(), date(int(m[1]), int(m[2]), int(m[3]))))
        except ValueError:
            pass
    for m in _NUMERIC.finditer(s):
        year = int(m[3]) + (2000 if len(m[3]) == 2 else 0)
        try:
            found.append((m.start(), date(year, int(m[1]), int(m[2]))))
        except ValueError:
            pass
    for m in _LONG.finditer(s):
        try:
            found.append((m.start(), date(int(m[3]), _month(m[1]), int(m[2]))))
        except ValueError:
            pass
    for m in _MONTH_YEAR.finditer(s):
        found.append((m.start(), date(int(m[2]), _month(m[1]), 1)))
    return min(found)[1] if found else None


def iso(value) -> str | None:
    d = parse_date(value)
    return d.isoformat() if d else None


def page_text(html: str) -> str:
    """Readable text of an HTML page, without scripts, styles and site chrome."""
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style", "noscript", "nav", "header", "footer", "form"]):
        tag.decompose()
    root = soup.find("main") or soup.find(id=re.compile("content|main", re.I)) or soup.body or soup
    lines = (line.strip() for line in root.get_text("\n").splitlines())
    return "\n".join(line for line in lines if line)


def field(text: str, label: str) -> str | None:
    """Value of a labeled field in page text: the line after the line equal to `label`."""
    lines = text.splitlines()
    for i, line in enumerate(lines[:-1]):
        if line.strip().lower() == label.lower():
            return lines[i + 1].strip() or None
    return None


def pdf_text(data: bytes, max_pages: int = 40) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = reader.pages[:max_pages]
        text = "\n".join((p.extract_text() or "") for p in pages)
        if len(reader.pages) > max_pages:
            text += f"\n[... {len(reader.pages) - max_pages} more pages not read]"
        return text
    except Exception as exc:  # malformed PDFs are common; never fail the run over one
        log.warning("Could not read PDF (%s)", exc)
        return ""


def sheet_rows(data: bytes) -> list[dict]:
    """Rows of the first worksheet of an .xlsx file, keyed by the header row."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    # The header is the first row with at least three non-empty cells.
    for i, row in enumerate(rows):
        if sum(1 for c in row if c not in (None, "")) >= 3:
            headers = [str(c).strip() if c is not None else f"col{j}" for j, c in enumerate(row)]
            out = []
            for r in rows[i + 1:]:
                if any(c not in (None, "") for c in r):
                    out.append({h: v for h, v in zip(headers, r)})
            return out
    return []


def clip(text: str, limit: int) -> str:
    return text if len(text) <= limit else text[:limit] + f"\n[... {len(text) - limit} more characters]"


def stable_id(*parts) -> str:
    """Short id from a record's identifying fields, for sources without record numbers."""
    raw = "|".join(re.sub(r"\s+", " ", str(p or "")).strip().lower() for p in parts)
    return hashlib.sha1(raw.encode()).hexdigest()[:12]
